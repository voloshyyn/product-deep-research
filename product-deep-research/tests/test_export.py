import copy
import json
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import load_workbook

import export_research
from export_research import export_files
from fixtures import dataset, write_dataset
from research_core import CATEGORIES, SCHEMAS, product_id


def rows(book, sheet):
    ws = book[sheet]
    columns = [c.value for c in ws[1]]
    return [dict(zip(columns, row)) for row in ws.iter_rows(min_row=2, values_only=True)]


def test_end_to_end_all_schemas_report_and_evidence(tmp_path):
    data = dataset(screenshot=True)
    source = write_dataset(tmp_path / "input", data)
    result = export_files([source], tmp_path / "output")
    run = result["runs"][0]
    assert all(Path(run[k]).exists() for k in ["report", "research_json", "screenshots", "manifest"])
    assert json.loads(Path(run["research_json"]).read_text()) == data
    doc = Document(run["report"])
    text = "\n".join(p.text for p in doc.paragraphs)
    assert all(names[0] in text for names in CATEGORIES.values())
    assert "Arbitrary integration criterion" in text
    assert all(c["statement"] in text for c in data["claims"])
    assert "PARTIALLY_VERIFIED" in text and "HYPOTHESIS" in text
    assert len(doc.inline_shapes) == 1
    with ZipFile(run["report"]) as archive:
        assert b"PAGE" in archive.read("word/footer1.xml")
        assert b"https://atlas.example/product" in archive.read("word/_rels/document.xml.rels")
    wb = load_workbook(result["database"])
    assert wb.sheetnames == list(SCHEMAS)
    for name, columns in SCHEMAS.items():
        assert [c.value for c in wb[name][1]] == columns
        assert wb[name].freeze_panes == "A2"
        assert wb[name].auto_filter.ref
    assert len(rows(wb, "Research History")) == 1
    assert len(rows(wb, "Custom Research")) == 1
    assert rows(wb, "UX Issues")[0]["severity"] == 2
    assert isinstance(rows(wb, "Research History")[0]["research_date"], datetime)
    assert any(c.hyperlink for row in wb["Sources"] for c in row)
    wb.close()


def test_factual_hypothesis_partial_and_company_separation(tmp_path):
    data = dataset()
    data["claims"][2].update(status="UNVERIFIED", confidence="LOW", evidence_ids=[], source_ids=[], gap_reason="Missing evidence")
    source = write_dataset(tmp_path / "in", data)
    result = export_files([source], tmp_path / "out")
    wb = load_workbook(result["database"])
    opportunity = rows(wb, "Opportunities")[0]
    assert opportunity["hypothesis"] and opportunity["confirmed_fact"] is None and opportunity["value"] is None
    company = next(r for r in rows(wb, "Overview") if r["claim_id"] == "c2")
    assert company["qualified_claim"] and company["confirmed_fact"] is None
    partial = next(r for r in rows(wb, "Pricing & Business") if r["claim_id"] == "c8")
    assert partial["qualified_claim"] and partial["qualification"] and partial["value"] is None
    price = next(r for r in rows(wb, "Pricing & Business") if r["claim_id"] == "c7")
    assert price["value"] == 12.5 and price["confirmed_fact"]
    unknown = next(r for r in rows(wb, "Audience & JTBD") if r["claim_id"] == "c3")
    assert unknown["gap"] and all(unknown[k] is None for k in ["confirmed_fact", "qualified_claim", "hypothesis", "value"])
    wb.close()


def test_idempotence_history_and_changed_data_collision(tmp_path):
    source = write_dataset(tmp_path / "in", dataset())
    root = tmp_path / "out"
    first = export_files([source], root)
    original_report = Path(first["runs"][0]["report"]).read_bytes()
    original_book = (root / "competitors.xlsx").read_bytes()
    retry = export_files([source], root)
    assert retry["runs"][0]["already_exported"]
    assert (root / "competitors.xlsx").read_bytes() == original_book
    data = dataset(run="run-002")
    source2 = write_dataset(tmp_path / "in2", data)
    second = export_files([source2], root)
    assert first["runs"][0]["report"] != second["runs"][0]["report"]
    assert Path(first["runs"][0]["report"]).read_bytes() == original_report
    wb = load_workbook(root / "competitors.xlsx")
    records = rows(wb, "Research History")
    assert len(records) == 2
    assert len({r["product_id"] for r in records}) == 1
    assert len({r["record_id"] for r in rows(wb, "Overview")}) == len(rows(wb, "Overview"))
    wb.close()
    assert len(list((root / ".backups").glob("*.xlsx"))) == 1
    data["claims"][0]["statement"] = "Changed evidence interpretation"
    write_dataset(source2.parent, data)
    with pytest.raises(ValueError, match="different data"):
        export_files([source2], root)


def test_multiple_industries_and_competitors_are_not_lost(tmp_path):
    inputs = [write_dataset(tmp_path / str(i), dataset(name=f"Synthetic {i}", url=f"https://p{i}.example",
                industry=industry)) for i, industry in enumerate(["healthcare", "finance", "e-commerce", "psychology", "productivity", "entertainment"])]
    result = export_files(inputs, tmp_path / "out")
    assert len(result["runs"]) == 6
    wb = load_workbook(result["database"])
    assert len(rows(wb, "Research History")) == 6
    assert len({r["product_id"] for r in rows(wb, "Research History")}) == 6
    assert len(rows(wb, "Custom Research")) == 6
    wb.close()


def test_arbitrary_custom_criteria_and_gap_get_records(tmp_path):
    data = dataset()
    data["research"]["additional_requirements"].append({"id": "custom-2", "title": "Особая проверка", "request": "Проверить сложное пользовательское условие"})
    data["custom_sections"].append({"requirement_id": "custom-2", "claim_ids": [], "limitation": "No public evidence"})
    source = write_dataset(tmp_path / "in", data)
    result = export_files([source], tmp_path / "out")
    wb = load_workbook(result["database"])
    assert {r["requirement_id"] for r in rows(wb, "Custom Research")} == {"custom-1", "custom-2"}
    assert rows(wb, "Custom Research")[1]["gap"] == "No public evidence"
    wb.close()


def test_russian_report(tmp_path):
    source = write_dataset(tmp_path / "in", dataset(language="ru"))
    result = export_files([source], tmp_path / "out")
    text = "\n".join(p.text for p in Document(result["runs"][0]["report"]).paragraphs)
    assert "Исследование продукта" in text and "Итоговые выводы" in text


def test_invalid_batch_creates_no_storage(tmp_path):
    first = write_dataset(tmp_path / "in1", dataset())
    bad = dataset(name="Bad", url="https://bad.example")
    bad["claims"][0]["evidence_ids"] = []
    second = write_dataset(tmp_path / "in2", bad)
    with pytest.raises(ValueError, match="validation failed"):
        export_files([first, second], tmp_path / "out")
    assert not (tmp_path / "out").exists()


def test_lock_and_schema_mismatch_preserve_files(tmp_path):
    source = write_dataset(tmp_path / "in", dataset())
    root = tmp_path / "out"
    root.mkdir()
    lock = root / ".export.lock"
    lock.write_text("another process")
    with pytest.raises(ValueError, match="locked"):
        export_files([source], root)
    assert lock.read_text() == "another process"
    lock.unlink()
    result = export_files([source], root)
    wb = load_workbook(result["database"])
    wb["Overview"]["A1"] = "unexpected_schema"
    wb.save(result["database"])
    wb.close()
    before = Path(result["database"]).read_bytes()
    with pytest.raises(ValueError, match="schema mismatch"):
        export_files([source], root)
    assert Path(result["database"]).read_bytes() == before


def test_rollback_on_publication_failure(tmp_path, monkeypatch):
    source = write_dataset(tmp_path / "in", dataset())
    root = tmp_path / "out"
    export_files([source], root)
    original = (root / "competitors.xlsx").read_bytes()
    source2 = write_dataset(tmp_path / "in2", dataset(run="run-002"))
    def fail_replace(*args):
        raise PermissionError("Workbook is open in Excel")
    monkeypatch.setattr(export_research.os, "replace", fail_replace)
    with pytest.raises(PermissionError, match="open in Excel"):
        export_files([source2], root)
    assert (root / "competitors.xlsx").read_bytes() == original
    assert not list(root.rglob("2025-01-15_run-002"))
    assert not (root / ".export.lock").exists()
    assert not list(root.glob(".export-*"))


def test_formula_injection_and_oversized_cells(tmp_path):
    data = dataset()
    data["product"]["name"] = '=HYPERLINK("https://malicious.example")'
    data["claims"][0]["statement"] = "=1+1"
    source = write_dataset(tmp_path / "in", data)
    result = export_files([source], tmp_path / "out")
    wb = load_workbook(result["database"])
    assert wb["Overview"]["C2"].data_type == "s"
    header = [c.value for c in wb["Overview"][1]]
    cell = wb["Overview"].cell(2, header.index("confirmed_fact") + 1)
    assert cell.value == "=1+1" and cell.data_type == "s"
    wb.close()
    data["research"]["run_id"] = "oversized"
    data["claims"][0]["statement"] = "x" * 32768
    source = write_dataset(tmp_path / "large", data)
    with pytest.raises(ValueError, match="no truncation"):
        export_files([source], tmp_path / "large-out")
    assert not list((tmp_path / "large-out").rglob("report.docx"))


def test_modified_artifacts_and_duplicate_batch_refused(tmp_path):
    source = write_dataset(tmp_path / "in", dataset())
    root = tmp_path / "out"
    with pytest.raises(ValueError, match="Duplicate product/run"):
        export_files([source, source], root)
    result = export_files([source], root)
    Path(result["runs"][0]["report"]).write_bytes(b"tampered")
    with pytest.raises(ValueError, match="missing or changed"):
        export_files([source], root)


def test_storage_environment_override_and_cli_precedence(tmp_path, monkeypatch):
    source = write_dataset(tmp_path / "in", dataset())
    monkeypatch.setenv("PRODUCT_RESEARCH_HOME", str(tmp_path / "configured"))
    assert Path(export_files([source])["database"]).parent == tmp_path / "configured"
    assert Path(export_files([source], tmp_path / "explicit")["database"]).parent == tmp_path / "explicit"


def test_strength_facet_keeps_category_and_synthesis_records(tmp_path):
    data = dataset()
    data["claims"][4]["facet"] = "strength"
    source = write_dataset(tmp_path / "in", data)
    result = export_files([source], tmp_path / "out")
    wb = load_workbook(result["database"])
    assert any(r["claim_id"] == "c5" for r in rows(wb, "Features & Services"))
    assert any(r["claim_id"] == "c5" for r in rows(wb, "Strengths & Weaknesses"))
    assert any(r["claim_id"] == "c15" for r in rows(wb, "Opportunities"))
    assert len(rows(wb, "Research History")) == 1
    wb.close()


def test_corrupt_manifest_and_changed_screenshot_refused(tmp_path):
    source = write_dataset(tmp_path / "in", dataset(screenshot=True))
    root = tmp_path / "out"
    result = export_files([source], root)
    manifest = Path(result["runs"][0]["manifest"])
    original = manifest.read_text()
    data = json.loads(original)
    data["files"] = {}
    manifest.write_text(json.dumps(data))
    with pytest.raises(ValueError, match="Incomplete artifact manifest"):
        export_files([source], root)
    manifest.write_text(original)
    # A PNG accepts trailing bytes but the integrity check must detect changed bytes.
    with (source.parent / "screenshots/synthetic.png").open("ab") as handle:
        handle.write(b"changed")
    with pytest.raises(ValueError, match="Screenshot changed"):
        export_files([source], root)
