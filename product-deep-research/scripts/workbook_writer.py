"""Append-only typed workbook records with exact schemas and safe string cells."""
from __future__ import annotations

import hashlib
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from research_core import CATEGORIES, COMMON, SCHEMAS, digest, normalize_url, product_id


def open_database(path: Path):
    if not path.exists():
        wb = Workbook()
        wb.remove(wb.active)
        for name, schema in SCHEMAS.items():
            wb.create_sheet(name).append(schema)
        wb.properties.title = "Product research evidence database"
        wb.properties.description = "Schema 1.0; append-only product/run records; identity is scope metadata"
        return wb
    wb = load_workbook(path)
    for name, schema in SCHEMAS.items():
        if name not in wb.sheetnames or [c.value for c in wb[name][1]] != schema:
            wb.close()
            raise ValueError(f"Workbook schema mismatch: {name}; existing workbook left unchanged")
        ids = [row[0] for row in wb[name].iter_rows(min_row=2, values_only=True)]
        if len(set(ids)) != len(ids) or any(not x for x in ids):
            wb.close()
            raise ValueError(f"Workbook contains duplicate or empty record IDs: {name}")
    return wb


def history_index(wb) -> dict[tuple[str, str], dict]:
    sheet = wb["Research History"]
    columns = [c.value for c in sheet[1]]
    records = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        record = dict(zip(columns, row))
        key = record["product_id"], record["run_id"]
        if key in records:
            raise ValueError(f"Duplicate research history: {key}")
        records[key] = record
    return records


def base_record(data: dict) -> dict:
    p, r = data["product"], data["research"]
    return dict(zip(COMMON[1:], [product_id(p), p["name"], p["url"], p["regional_variant"],
        p["industry"], p["market"], r["run_id"], date.fromisoformat(r["date"]), r["comparison_group"]]))


def claim_record(claim: dict, data: dict) -> dict:
    sources = {s["id"]: s for s in data["sources"]}
    status = claim["status"]
    confirmed = status == "VERIFIED" and claim["kind"] in {"FACT", "OBSERVATION"}
    qualified = status == "PARTIALLY_VERIFIED" or (status == "VERIFIED" and claim["kind"] == "COMPANY_CLAIM")
    return {
        "claim_id": claim["id"], "category": claim["category"], "title": claim["title"],
        "claim_kind": claim["kind"], "facet": claim["facet"],
        "confirmed_fact": claim["statement"] if confirmed else None,
        "qualified_claim": claim["statement"] if qualified else None,
        "hypothesis": claim["statement"] if status == "HYPOTHESIS" else None,
        "gap": f"{claim['statement']} | {claim['gap_reason']}" if status in {"UNVERIFIED", "NOT_ACCESSIBLE"} else None,
        # Values are strictly confirmed; partial/company/hypothetical values remain in JSON/report.
        "value": claim["value"] if confirmed else None, "unit": claim["unit"] if confirmed else None,
        "qualification": claim["qualification"], "verification_status": status,
        "confidence": claim["confidence"], "verified_on": date.fromisoformat(claim["verified_on"]),
        "evidence_ids": "; ".join(claim["evidence_ids"]),
        "source_urls": "\n".join(sources[s]["url"] for s in claim["source_ids"]),
    }


def append_run(wb, data: dict, report: Path, dataset: Path) -> None:
    base = base_record(data)
    reqs = {r["id"]: r for r in data["research"]["additional_requirements"]}
    claims = {c["id"]: c for c in data["claims"]}

    def append(sheet: str, local_id: str, values: dict) -> None:
        key = f"{base['product_id']}:{base['run_id']}:{sheet}:{local_id}"
        rid = "rec_" + hashlib.sha256(key.encode()).hexdigest()[:32]
        record = {**base, **values, "record_id": rid}
        row_num = wb[sheet].max_row + 1
        for col, name in enumerate(SCHEMAS[sheet], 1):
            value = record.get(name)
            cell = wb[sheet].cell(row_num, col, value)
            if isinstance(value, str):
                if len(value) > 32767:
                    raise ValueError(f"Excel cell exceeds 32767 characters: {sheet}/{local_id}/{name}; no truncation performed")
                # Prevent formulas and external-link injection, including strings beginning '='.
                cell.data_type = "s"
                if name in {"product_url", "source_url", "page_url", "source_urls"}:
                    url = value.splitlines()[0] if value else ""
                    if url.startswith(("https://", "http://")):
                        cell.hyperlink = url
                        cell.style = "Hyperlink"

    for c in claims.values():
        if c["category"] == "custom":
            req = reqs[c["custom_requirement_id"]]
            append("Custom Research", c["id"], {**claim_record(c, data), "requirement_id": req["id"],
                "requirement_title": req["title"], "user_request": req["request"]})
        else:
            sheet = CATEGORIES[c["category"]][2]
            append(sheet, c["id"], claim_record(c, data))
        facet_sheet = {"strength": "Strengths & Weaknesses", "weakness": "Strengths & Weaknesses", "opportunity": "Opportunities"}.get(c["facet"])
        primary_sheet = "Custom Research" if c["category"] == "custom" else CATEGORIES[c["category"]][2]
        if facet_sheet and facet_sheet != primary_sheet:
            append(facet_sheet, c["id"], claim_record(c, data))
    for section in [*data["sections"], *data["custom_sections"]]:
        if not section["limitation"]:
            continue
        key = section.get("category", "custom")
        values = {"category": key, "title": "Coverage limitation", "gap": section["limitation"],
                  "verification_status": "UNVERIFIED", "confidence": "LOW",
                  "verified_on": date.fromisoformat(data["research"]["date"])}
        if key == "custom":
            req = reqs[section["requirement_id"]]
            values.update(requirement_id=req["id"], requirement_title=req["title"], user_request=req["request"])
            append("Custom Research", "gap-" + req["id"], values)
        else:
            append(CATEGORIES[key][2], "gap-" + key, values)
    for source in data["sources"]:
        evs = [e for e in data["evidence"] if e["source_id"] == source["id"]]
        for ev in evs or [None]:
            linked = [c for c in claims.values() if source["id"] in c["source_ids"] and
                      (ev is None or ev["id"] in c["evidence_ids"])]
            values = {"source_id": source["id"], "source_url": source["url"], "source_title": source["title"],
                "publisher_type": source["publisher_type"], "independence_group": source["independence_group"],
                "access_status": source["access_status"], "accessed_on": date.fromisoformat(source["accessed_on"]),
                "access_method": source["access_method"], "access_detail": source["access_detail"],
                "verification_status": "; ".join(sorted({c["status"] for c in linked})) or "UNVERIFIED",
                "confidence": "; ".join(sorted({c["confidence"] for c in linked})) or "LOW",
                "linked_claim_ids": "; ".join(c["id"] for c in linked)}
            if ev:
                values.update(evidence_id=ev["id"], evidence_kind=ev["kind"], verified_on=date.fromisoformat(ev["verified_on"]),
                    evidence_detail=ev["detail"], locator=ev["locator"], screenshot=ev["screenshot"])
            append("Sources", ev["id"] if ev else source["id"], values)
    for issue in data["ux_issues"]:
        c = claims[issue["claim_id"]]
        cr = claim_record(c, data)
        values = {**issue, "issue_id": issue["id"], "heuristics": "; ".join(map(str, issue["heuristics"])),
            "wcag": "; ".join(issue["wcag"]), "verification_status": c["status"], "confidence": c["confidence"],
            "evidence_ids": cr["evidence_ids"], "source_urls": cr["source_urls"]}
        append("UX Issues", issue["id"], values)
    append("Research History", "run", {
        "canonical_url": normalize_url(data["product"]["url"]), "dataset_sha256": digest(data),
        "report_path": str(report), "dataset_path": str(dataset), "language": data["research"]["language"],
        "depth": data["research"]["depth"], "objectives": "\n".join(data["research"]["objectives"]),
        "methodology": data["research"]["methodology"], "limitations": "\n".join(data["limitations"]),
        "claim_count": len(claims), "gap_count": sum(c["status"] in {"UNVERIFIED", "NOT_ACCESSIBLE"} for c in claims.values()),
        "verification_status": "; ".join(sorted({c["status"] for c in claims.values()})) or "UNVERIFIED",
        "evidence_ids": "; ".join(e["id"] for e in data["evidence"]),
        "source_urls": "\n".join(s["url"] for s in data["sources"]),
    })


def format_database(wb) -> None:
    for index, (name, schema) in enumerate(SCHEMAS.items(), 1):
        ws = wb[name]
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        ws.print_title_rows = "1:1"
        ws.row_dimensions[1].height = 32
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="19354D")
            cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for i, column in enumerate(schema, 1):
            width = 48 if column in {"confirmed_fact", "qualified_claim", "hypothesis", "gap", "evidence_detail", "methodology", "limitations", "user_request"} else 22
            if "url" in column or "path" in column:
                width = 48
            ws.column_dimensions[get_column_letter(i)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"
                if not cell.hyperlink:
                    cell.font = Font(name="Calibri", size=10)
        ref = f"A1:{get_column_letter(len(schema))}{max(ws.max_row, 1)}"
        ws.auto_filter.ref = ref
        table_name = f"ResearchTable{index:02d}"
        if table_name in ws.tables:
            ws.tables[table_name].ref = ref
        elif ws.max_row > 1:
            table = Table(displayName=table_name, ref=ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
