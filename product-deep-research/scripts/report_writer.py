"""Readable, evidence-linked python-docx reports; content comes only from the dataset."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from research_core import CATEGORIES, product_id, safe_screenshot

LABELS = {
    "en": {
        "report": "Product research", "summary": "Executive summary", "scope": "Objectives scope and methodology",
        "custom": "Additional research requirements", "journeys": "User journey testing", "issues": "Detailed UX findings",
        "recommendations": "Prioritized recommendations", "verification": "Evidence and verification",
        "sources": "Source references", "limitations": "Research limitations", "conclusions": "Final conclusions",
        "industry": "Industry", "market": "Target market", "date": "Research date", "method": "Methodology",
        "depth": "Depth", "language": "Language", "context": "User supplied identity and research context",
        "evidence": "Evidence", "qualification": "Qualification and reasoning", "gap": "Verification limitation",
        "value": "Structured value", "none": "No records supplied", "status": "Status", "claim": "Finding",
        "test": "Testing", "goal": "Goal", "severity": "Severity", "impact": "Expected user impact",
        "solution": "Recommended solution", "rationale": "Rationale", "request": "User requirement",
        "independence": "Publisher and independence group", "access": "Access method and result",
        "classification": "Classification", "confidence": "Confidence", "verified": "Verified on",
        "screenshots": "Evidence screenshots", "comparison": "Competitor comparison scope",
        "verification_note": "Statuses describe evidence support. COMPANY_CLAIM verifies attribution only. HYPOTHESIS remains an interpretation. Severity 0–4 and recommendation priorities are analyst judgments. Source independence requires editorial review; repeated unsupported statements are not corroboration.",
    },
    "ru": {
        "report": "Исследование продукта", "summary": "Краткие выводы", "scope": "Цели охват и методология",
        "custom": "Дополнительные требования исследования", "journeys": "Проверка пользовательских сценариев", "issues": "Подробные UX находки",
        "recommendations": "Приоритетные рекомендации", "verification": "Доказательства и верификация",
        "sources": "Источники", "limitations": "Ограничения исследования", "conclusions": "Итоговые выводы",
        "industry": "Отрасль", "market": "Целевой рынок", "date": "Дата исследования", "method": "Методология",
        "depth": "Глубина", "language": "Язык", "context": "Идентификация и контекст исследования по запросу пользователя",
        "evidence": "Доказательства", "qualification": "Оговорки и обоснование", "gap": "Ограничение проверки",
        "value": "Структурированное значение", "none": "Записи не предоставлены", "status": "Статус", "claim": "Находка",
        "test": "Проверка", "goal": "Цель", "severity": "Серьезность", "impact": "Ожидаемое влияние на пользователя",
        "solution": "Предлагаемое решение", "rationale": "Обоснование", "request": "Требование пользователя",
        "independence": "Издатель и группа независимости", "access": "Способ и результат доступа",
        "classification": "Классификация", "confidence": "Уверенность", "verified": "Дата проверки",
        "screenshots": "Скриншоты доказательств", "comparison": "Область сравнения конкурентов",
        "verification_note": "Статусы отражают достаточность доказательств. COMPANY_CLAIM подтверждает только атрибуцию заявления. HYPOTHESIS остается интерпретацией. Серьезность 0–4 и приоритеты рекомендаций являются оценками исследователя. Независимость источников требует содержательной проверки. Повторение неподтвержденных заявлений не является верификацией.",
    },
}


def labels_for(data: dict) -> dict[str, str]:
    lang = data["research"]["language"].split("-")[0]
    labels = dict(LABELS.get(lang, LABELS["en"]))
    labels.update(data.get("labels", {}))
    for key, names in CATEGORIES.items():
        labels.setdefault(key, names[1 if lang == "ru" else 0])
    return labels


def hyperlink(paragraph, url: str) -> None:
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    properties.append(style)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = url
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def table(doc, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    tab = doc.add_table(rows=1, cols=len(headers))
    tab.autofit = False
    props = tab._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
            border.set(qn(f"w:{key}"), value)
        borders.append(border)
    props.append(borders)
    repeat = OxmlElement("w:tblHeader")
    tab.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_index, values in enumerate([headers, *rows]):
        cells = tab.rows[0].cells if row_index == 0 else tab.add_row().cells
        for i, value in enumerate(values):
            cells[i].width = Inches(widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[i].text = str(value)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E6EDF5" if row_index == 0 else ("F7F9FC" if row_index % 2 == 0 else "FFFFFF"))
            cells[i]._tc.get_or_add_tcPr().append(shading)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement(f"w:{edge}")
                margin.set(qn("w:w"), "100")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            cells[i]._tc.get_or_add_tcPr().append(margins)
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.space_before = Pt(4)
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.bold = row_index == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def write_report(data: dict, output: Path, base_dir: Path) -> None:
    labels = labels_for(data)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.85)
    section.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in [("Title", 30), ("Subtitle", 13), ("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 11)]:
        style = doc.styles[name]
        style.font.name, style.font.size, style.font.color.rgb = "Calibri", Pt(size), RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(7)
    doc.core_properties.title = f"{labels['report']} {data['product']['name']}"
    doc.core_properties.subject = product_id(data["product"])
    doc.core_properties.language = data["research"]["language"]
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    p = data["product"]
    r = data["research"]
    doc.add_paragraph(labels["report"], "Title")
    doc.add_paragraph(p["name"], "Subtitle")
    doc.add_paragraph(labels["context"])
    for label, value in [("industry", p["industry"]), ("market", p["market"]), ("date", r["date"]), ("depth", r["depth"]), ("language", r["language"])]:
        doc.add_paragraph(f"{labels[label]}: {value}")
    hyperlink(doc.add_paragraph(), p["url"])
    if p["regional_variant"]:
        doc.add_paragraph(p["regional_variant"])
    doc.add_paragraph(f"{product_id(p)} / {r['run_id']}")
    doc.add_page_break()
    claims = {c["id"]: c for c in data["claims"]}
    evidence = {e["id"]: e for e in data["evidence"]}
    sources = {s["id"]: s for s in data["sources"]}

    def render_claim(cid: str, heading: bool = True) -> None:
        c = claims[cid]
        if heading:
            doc.add_heading(c["title"], 2)
        doc.add_paragraph(f"[{cid}] {c['statement']}")
        meta = f"{labels['classification']}: {c['kind']} / {c['status']} · {labels['confidence']}: {c['confidence']} · {labels['verified']}: {c['verified_on']}"
        paragraph = doc.add_paragraph(meta)
        for run in paragraph.runs:
            run.font.size = Pt(9)
        for key, label in [("qualification", "qualification"), ("gap_reason", "gap")]:
            if c[key]:
                doc.add_paragraph(f"{labels[label]}: {c[key]}")
        if c["value"] is not None:
            doc.add_paragraph(f"{labels['value']}: {c['value']} {c['unit']}")
        if c["evidence_ids"]:
            doc.add_paragraph(f"{labels['evidence']}: {', '.join(c['evidence_ids'])}")
        for sid in c["source_ids"]:
            paragraph = doc.add_paragraph(f"[{sid}] ")
            hyperlink(paragraph, sources[sid]["url"])

    def render_refs(ids: list[str]) -> None:
        if not ids:
            doc.add_paragraph(labels["none"])
        for cid in ids:
            render_claim(cid, False)

    doc.add_heading(labels["summary"], 1)
    render_refs(data["executive_summary"])
    doc.add_heading(labels["scope"], 1)
    for objective in r["objectives"]:
        doc.add_paragraph(objective, "List Bullet")
    doc.add_paragraph(f"{labels['method']}: {r['methodology']}")
    if r["competitor_urls"] or r["comparison_group"]:
        doc.add_heading(labels["comparison"], 2)
        doc.add_paragraph(r["comparison_group"])
        for url in r["competitor_urls"]:
            hyperlink(doc.add_paragraph(), url)
    for section_data in data["sections"]:
        key = section_data["category"]
        doc.add_heading(labels[key], 1)
        for cid in section_data["claim_ids"]:
            render_claim(cid)
        if section_data["limitation"]:
            doc.add_paragraph(f"{labels['gap']}: {section_data['limitation']}")
        if key == "pricing" and section_data["claim_ids"]:
            rows = [[c["title"], str(c["value"]) + " " + c["unit"] if c["value"] is not None else "—", c["status"], ", ".join(c["evidence_ids"])]
                    for c in (claims[cid] for cid in section_data["claim_ids"])]
            table(doc, [labels["claim"], labels["value"], labels["status"], labels["evidence"]], rows, [2.1, 1.2, 1.55, 1.7])
    reqs = {item["id"]: item for item in r["additional_requirements"]}
    if reqs:
        doc.add_heading(labels["custom"], 1)
    for section_data in data["custom_sections"]:
        req = reqs[section_data["requirement_id"]]
        doc.add_heading(req["title"], 2)
        doc.add_paragraph(f"{labels['request']}: {req['request']}")
        for cid in section_data["claim_ids"]:
            render_claim(cid)
        if section_data["limitation"]:
            doc.add_paragraph(f"{labels['gap']}: {section_data['limitation']}")
    doc.add_heading(labels["journeys"], 1)
    for journey in data["journeys"]:
        doc.add_heading(journey["title"], 2)
        doc.add_paragraph(f"{labels['goal']}: {journey['goal']}")
        doc.add_paragraph(f"{labels['test']}: {journey['status']} · {labels['evidence']}: {', '.join(journey['evidence_ids'])}")
        for step in journey["steps"]:
            doc.add_paragraph(step["action"], "List Number")
            hyperlink(doc.add_paragraph(), step["url"])
            render_refs(step["outcome_claim_ids"])
        if journey["limitation"]:
            doc.add_paragraph(f"{labels['gap']}: {journey['limitation']}")
    doc.add_heading(labels["issues"], 1)
    if not data["ux_issues"]:
        doc.add_paragraph(labels["none"])
    for issue in sorted(data["ux_issues"], key=lambda i: (-i["severity"], i["id"])):
        doc.add_heading(issue["title"], 2)
        doc.add_paragraph(f"{issue['id']} · {labels['severity']}: {issue['severity']}/4 · {issue['priority']} · Nielsen {', '.join(map(str, issue['heuristics']))}")
        hyperlink(doc.add_paragraph(), issue["page_url"])
        render_claim(issue["claim_id"], False)
        for key, label in [("severity_rationale", "rationale"), ("impact", "impact"), ("recommendation", "solution")]:
            doc.add_paragraph(f"{labels[label]}: {issue[key]}")
        if issue["wcag"]:
            doc.add_paragraph("WCAG: " + ", ".join(issue["wcag"]))
    doc.add_heading(labels["recommendations"], 1)
    for rec in sorted(data["recommendations"], key=lambda x: (x["priority"], x["id"])):
        doc.add_heading(f"{rec['priority']} {rec['title']}", 2)
        doc.add_paragraph(rec["action"])
        doc.add_paragraph(f"{labels['rationale']}: {rec['rationale']} [{', '.join(rec['claim_ids'])}]")
    doc.add_heading(labels["verification"], 1)
    doc.add_paragraph(labels["verification_note"])
    for ev in evidence.values():
        doc.add_heading(ev["id"], 2)
        doc.add_paragraph(f"{ev['kind']} · {ev['verified_on']} · {ev['source_id']}")
        doc.add_paragraph(ev["detail"])
        doc.add_paragraph(ev["locator"])
        hyperlink(doc.add_paragraph(), sources[ev["source_id"]]["url"])
        if ev["screenshot"]:
            image = safe_screenshot(base_dir, ev["screenshot"])
            from docx.image.image import Image
            picture = Image.from_file(str(image))
            width = min(6.4, 6.4 * picture.px_width / picture.px_height)
            doc.add_picture(str(image), width=Inches(width))
            doc.add_paragraph(f"{labels['screenshots']}: {ev['id']} · {ev['screenshot']}", "Caption")
    doc.add_heading(labels["sources"], 1)
    for s in sources.values():
        doc.add_heading(f"{s['id']} {s['title']}", 2)
        hyperlink(doc.add_paragraph(), s["url"])
        doc.add_paragraph(f"{labels['independence']}: {s['publisher_type']} / {s['independence_group']}")
        doc.add_paragraph(f"{labels['access']}: {s['accessed_on']} / {s['access_method']} / {s['access_status']}")
        doc.add_paragraph(s["access_detail"])
    doc.add_heading(labels["limitations"], 1)
    for limitation in data["limitations"]:
        doc.add_paragraph(limitation, "List Bullet")
    for c in claims.values():
        if c["status"] in {"UNVERIFIED", "NOT_ACCESSIBLE", "PARTIALLY_VERIFIED"}:
            doc.add_paragraph(f"[{c['id']}] {c['status']}: {c['gap_reason'] or c['qualification']}")
    for s in [*data["sections"], *data["custom_sections"]]:
        if s["limitation"]:
            doc.add_paragraph(f"{s.get('category', s.get('requirement_id'))}: {s['limitation']}")
    doc.add_heading(labels["conclusions"], 1)
    render_refs(data["conclusions"])
    doc.save(output)
